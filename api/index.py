import os
import io
import json
import re

import requests
from fastapi import FastAPI, APIRouter, HTTPException, Request
from fastapi.responses import Response, JSONResponse
from pydantic import BaseModel
from groq import Groq
from docx import Document
from docx.shared import Pt, RGBColor

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

app = FastAPI()
router = APIRouter()


@app.exception_handler(404)
async def not_found(request: Request, exc):
    # Shows which path Vercel actually gave us, so we can debug routing.
    return JSONResponse({"detail": "Not Found", "saw_path": request.url.path}, status_code=404)

# Don't crash the whole site at startup if the key is missing on the server.
GROQ_KEY = os.environ.get("GROQ_API_KEY")
client = Groq(api_key=GROQ_KEY) if GROQ_KEY else None
MODEL = "openai/gpt-oss-120b"

PERSONALITY = """You are 'Study Bestie', a sweet, cheerful, girly study buddy 🎀.
- Talk like a supportive best friend: "hey bestie!", "okay so", "you've got this!".
- Use cute emojis lightly (🎀 🌸 💖 ✨ 🩷), 2 to 4 per reply, never spammy.
- Stay accurate and clear. Being cute must never make the explanation worse.
- Explain simply, like to a friend before an exam. End with one short encouraging line."""

CHAT_SYSTEM = PERSONALITY + """

You get excerpts from a PDF and a question. Answer using ONLY the excerpts.
If the answer isn't there, say kindly you couldn't find it in the PDF 🥺 and suggest rephrasing.
Use short paragraphs or bullets when it helps."""

SUMMARY_SYSTEM = PERSONALITY + """

Summarize the PDF text. Start with a friendly one-line greeting and one sentence on what
the PDF is about, then list the MAIN topics and key points as short bullets."""


class Summary(BaseModel):
    text: str


class Chat(BaseModel):
    question: str
    context: str
    history: list[dict] = []


class Doc(BaseModel):
    text: str


def generate(messages):
    if client is None:
        raise HTTPException(status_code=500, detail="GROQ_API_KEY is not set on the server")
    try:
        r = client.chat.completions.create(model=MODEL, messages=messages)
        return r.choices[0].message.content
    except Exception as e:
        code = 429 if "429" in str(e) else 500
        raise HTTPException(status_code=code, detail=str(e))


def gen_json(task, text):
    raw = generate([
        {"role": "system", "content": "Reply with ONLY valid JSON, no extra text."},
        {"role": "user", "content": f"{task}\n\nStudy material:\n{text[:12000]}"},
    ])
    try:
        return json.loads(re.search(r"[\[{].*[\]}]", raw, re.S).group(0))
    except Exception:
        raise HTTPException(status_code=500, detail="Bad JSON from model")


def as_list(data):
    """Sometimes the model wraps the list in an object. Pull the list out."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list):
                return v
    return []


@router.post("/summary")
def summary(body: Summary):
    return {"reply": generate([
        {"role": "system", "content": SUMMARY_SYSTEM},
        {"role": "user", "content": f"PDF text:\n{body.text[:14000]}"},
    ])}


@router.post("/chat")
def chat(body: Chat):
    messages = [{"role": "system", "content": CHAT_SYSTEM}]
    for m in body.history[-6:]:
        messages.append({"role": m["role"], "content": m["content"]})
    messages.append({"role": "user",
                     "content": f"PDF excerpts:\n{body.context}\n\nQuestion: {body.question}"})
    return {"reply": generate(messages)}


@router.post("/flashcards")
def flashcards(body: Doc):
    return {"items": as_list(gen_json(
        'Make 10 flashcards. JSON: [{"q":"question","a":"short answer"}]', body.text))}


@router.post("/quiz")
def quiz(body: Doc):
    return {"items": as_list(gen_json(
        'Make 5 multiple-choice questions. JSON: [{"q":"...","options":["...","...","...","..."],'
        '"answer":0,"why":"one-line explanation"}] where answer is the index of the right option.',
        body.text))}


@router.post("/shortforms")
def shortforms(body: Doc):
    return {"items": as_list(gen_json(
        'Find the abbreviations and short forms in the material (max 12) and give each a fun memory '
        'trick. JSON: [{"short":"BFS","full":"Breadth First Search","trick":"..."}]. '
        'If there are none, return [].', body.text))}


@router.post("/notes")
def notes(body: Doc):
    n = gen_json(
        'Make exam notes. JSON: {"title":"...","overview":"...","sections":'
        '[{"heading":"...","points":["..."]}],"remember":["..."]}', body.text)
    if not isinstance(n, dict):
        raise HTTPException(status_code=500, detail="Bad notes format")
    doc = Document()
    pink = RGBColor(0xE9, 0x4F, 0x8A)

    def head(t, size=15):
        r = doc.add_paragraph().add_run(t)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = pink

    head("🎀 " + n.get("title", "Study Notes"), 24)
    doc.add_paragraph(n.get("overview", ""))
    for s in n.get("sections", []):
        head("✨ " + s.get("heading", ""))
        for p in s.get("points", []):
            doc.add_paragraph(str(p), style="List Bullet")
    if n.get("remember"):
        head("💡 Remember This")
        for p in n["remember"]:
            doc.add_paragraph(str(p), style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@router.post("/videos")
def videos(body: Doc):
    yt_key = os.environ.get("YOUTUBE_API_KEY")
    if not yt_key:
        raise HTTPException(status_code=500, detail="YOUTUBE_API_KEY is not set on the server")

    queries = as_list(gen_json(
        'Give 3 short YouTube search queries that would help a student understand the main '
        'concepts of this material. JSON: ["query one","query two","query three"]', body.text))

    items, seen = [], set()
    for q in queries[:3]:
        try:
            r = requests.get(
                "https://www.googleapis.com/youtube/v3/search",
                params={
                    "part": "snippet",
                    "q": str(q) + " explained",
                    "type": "video",
                    "maxResults": 2,
                    "videoEmbeddable": "true",
                    "safeSearch": "strict",
                    "key": yt_key,
                },
                timeout=10,
            )
        except requests.RequestException:
            continue
        if r.status_code == 403:
            raise HTTPException(status_code=429, detail="YouTube quota reached or API not enabled")
        if r.status_code != 200:
            continue
        for v in r.json().get("items", []):
            vid = v.get("id", {}).get("videoId")
            if not vid or vid in seen:
                continue
            seen.add(vid)
            sn = v["snippet"]
            items.append({
                "id": vid,
                "title": sn["title"],
                "channel": sn["channelTitle"],
                "thumb": sn.get("thumbnails", {}).get("medium", {}).get("url", ""),
            })
    return {"items": items}


@router.get("/health")
def health():
    return {
        "ok": True,
        "groq_key": bool(GROQ_KEY),
        "youtube_key": bool(os.environ.get("YOUTUBE_API_KEY")),
    }


# Work whether Vercel passes /api/summary, /summary or /api/index/summary
app.include_router(router, prefix="/api")
app.include_router(router, prefix="/api/index")
app.include_router(router)


# Local development only: serve the website from the same server.
# Keep this LAST, after all the routes above.
if not os.environ.get("VERCEL"):
    from fastapi.staticfiles import StaticFiles
    folder = os.path.join(os.path.dirname(__file__), "..", "public")
    app.mount("/", StaticFiles(directory=folder, html=True), name="site")